"""CT-FAT-001..003 do Django — spec relatorio-faturas-pendentes."""
from io import BytesIO
from unittest.mock import patch

import requests
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from fedhub import relatorios_documentos as documentos
from fedhub.services.relatorios_service import (
    FedHubDemorou,
    FedHubIndisponivel,
    FedHubRecusou,
    RelatoriosFedhubService,
    filtrar_parametros,
)

Usuario = get_user_model()

URL = "/faturas/pendentes/"
URL_XLSX = "/faturas/pendentes/exportar-excel/"
URL_PDF = "/faturas/pendentes/exportar-pdf/"


def linha(**sobrescreve):
    base = {
        "fatura": 169212, "tipo_fat": "CONDOCORP", "documento": "0001510980", "sacado": "CONDOMINIO EXEMPLO",
        "cnpj_sacado": "00000000000191", "produto": "PROGRAMA DE GESTÃO", "obs": "", "vigencia": "05/2026",
        "dt_ini_vig": "2026-05-01", "vencimento": "2026-05-20", "dias_atraso": 112, "valor": 190.80,
        "parcela": "4/1", "administradora": "5", "administradora_nome": "ADMINISTRADORA EXEMPLO LTDA",
        "seguradora": "10", "seguradora_nome": "SEGURADORA X", "cedente": "3", "cedente_cnpj": "00000000000272",
        "apolice": "AP1", "nosso_numero": "123", "deposito_cc": "", "tem_deposito_cc": False,
    }
    base.update(sobrescreve)
    return base


def resposta_fedhub(linhas=None):
    linhas = linhas if linhas is not None else [linha(), linha(documento="0001510981", valor=213.80)]
    total = round(sum(l["valor"] for l in linhas), 2)
    return {
        "status": "success",
        "filtros": {"situacao": "vencidas", "classificacao": "sem_obs_sem_deposito", "ordenar_por": "vencimento"},
        "referencia": "2026-09-09",
        "totais": {
            "documentos": len(linhas), "faturas": len({l["fatura"] for l in linhas}),
            "valor_total": total, "valor_pago": 0.0,
            "por_administradora": [{"administradora": "5", "nome": "ADMINISTRADORA EXEMPLO LTDA",
                                    "documentos": len(linhas), "valor": total}] if linhas else [],
        },
        "data": linhas,
        "timestamp": "2026-09-09T10:00:00",
    }


class _Base(APITestCase):
    def setUp(self):
        self.financeiro = Usuario.objects.create_user(email="fin@t.com", password="x", nivel_acesso="financeiro")
        self.faturamento = Usuario.objects.create_user(email="fat@t.com", password="x", nivel_acesso="faturamento")
        self.admin = Usuario.objects.create_user(email="adm@t.com", password="x", nivel_acesso="admin")
        self.comum = Usuario.objects.create_user(email="usr@t.com", password="x", nivel_acesso="usuario")
        self.client.force_authenticate(self.financeiro)


@patch.object(RelatoriosFedhubService, "faturas_pendentes")
class FaturasPendentesDadosTests(_Base):
    """CT-FAT-001: proxy, filtros, erros do FedHub e permissões."""

    def test_repassa_filtros_e_devolve_linhas_e_totais_sem_recontar(self, buscar):
        buscar.return_value = resposta_fedhub()

        resposta = self.client.get(URL, {"administradora": "5", "situacao": "todas", "ordenar_por": "fatura", "vazio": ""})

        self.assertEqual(resposta.status_code, status.HTTP_200_OK, resposta.data)
        self.assertTrue(resposta.data["sucesso"])
        self.assertEqual(resposta.data["totais"]["documentos"], 2)
        self.assertEqual(resposta.data["totais"]["valor_total"], 404.60)
        self.assertEqual(len(resposta.data["data"]), 2)
        buscar.assert_called_once_with({"administradora": "5", "situacao": "todas", "ordenar_por": "fatura"})

    def test_parametros_fora_do_contrato_nao_sao_repassados(self, buscar):
        buscar.return_value = resposta_fedhub([])

        self.client.get(URL, {"produtor": "1", "pagamento_ini": "2026-01-01", "fatura": "10", "page": "2"})

        buscar.assert_called_once_with({"fatura": "10"})

    def test_fedhub_fora_502_e_demorou_504(self, buscar):
        buscar.side_effect = FedHubIndisponivel()
        fora = self.client.get(URL)
        buscar.side_effect = FedHubDemorou()
        demorou = self.client.get(URL)

        self.assertEqual(fora.status_code, status.HTTP_502_BAD_GATEWAY)
        self.assertEqual(demorou.status_code, status.HTTP_504_GATEWAY_TIMEOUT)
        self.assertIn("erro", fora.data)

    def test_filtro_invalido_no_fedhub_volta_400(self, buscar):
        buscar.side_effect = FedHubRecusou(400, {"detail": "situacao deve ser um de vencidas, a_vencer, todas."})

        resposta = self.client.get(URL, {"situacao": "atrasadas"})

        self.assertEqual(resposta.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("situacao", str(resposta.data["erro"]))

    def test_outro_erro_do_fedhub_volta_502(self, buscar):
        buscar.side_effect = FedHubRecusou(503, {"detail": "Firebird indisponível"})

        resposta = self.client.get(URL)

        self.assertEqual(resposta.status_code, status.HTTP_502_BAD_GATEWAY)

    def test_permissoes_por_nivel(self, buscar):
        buscar.return_value = resposta_fedhub([])
        esperado = {self.financeiro: 200, self.faturamento: 200, self.admin: 200, self.comum: 403}

        for usuario, codigo in esperado.items():
            self.client.force_authenticate(usuario)
            for url in (URL, URL_XLSX, URL_PDF):
                self.assertEqual(self.client.get(url).status_code, codigo, (usuario.nivel_acesso, url))

    def test_sem_login_401(self, buscar):
        self.client.force_authenticate(None)

        self.assertEqual(self.client.get(URL).status_code, status.HTTP_401_UNAUTHORIZED)


@patch.object(RelatoriosFedhubService, "faturas_pendentes")
class FaturasPendentesExcelTests(_Base):
    """CT-FAT-002: planilha abre, colunas e totais batem, vazio devolve planilha vazia."""

    def test_planilha_uma_linha_por_documento_com_totais(self, buscar):
        buscar.return_value = resposta_fedhub()

        resposta = self.client.get(URL_XLSX, {"situacao": "vencidas"})

        self.assertEqual(resposta.status_code, status.HTTP_200_OK)
        self.assertIn("faturas-pendentes-2026-09-09.xlsx", resposta["Content-Disposition"])
        self.assertEqual(resposta["Access-Control-Expose-Headers"], "Content-Disposition")
        wb = load_workbook(BytesIO(resposta.content))
        ws = wb["Faturas pendentes"]
        cabecalho = [c.value for c in ws[1]]
        self.assertEqual(cabecalho[:3], ["Fatura", "Documento", "Sacado"])
        self.assertEqual(ws.cell(row=2, column=2).value, "0001510980")
        self.assertEqual(ws.cell(row=2, column=9).value, 190.80)  # número, não texto
        self.assertEqual(ws.cell(row=2, column=7).value.strftime("%Y-%m-%d"), "2026-05-20")  # data
        self.assertEqual(ws.cell(row=4, column=8).value, "TOTAL GERAL")
        self.assertEqual(ws.cell(row=4, column=9).value, 404.60)
        self.assertIn("Filtros", wb.sheetnames)

    def test_planilha_vazia_tem_cabecalho_e_totais_zerados(self, buscar):
        buscar.return_value = resposta_fedhub([])

        resposta = self.client.get(URL_XLSX)

        self.assertEqual(resposta.status_code, status.HTTP_200_OK)
        ws = load_workbook(BytesIO(resposta.content))["Faturas pendentes"]
        self.assertEqual(ws.cell(row=1, column=1).value, "Fatura")
        self.assertEqual(ws.cell(row=2, column=8).value, "TOTAL GERAL")
        self.assertEqual(ws.cell(row=2, column=9).value, 0.0)


@patch.object(RelatoriosFedhubService, "faturas_pendentes")
class FaturasPendentesPDFTests(_Base):
    """CT-FAT-003: PDF com o layout do legado."""

    def test_pdf_e_gerado_com_nome_e_header_expostos(self, buscar):
        buscar.return_value = resposta_fedhub()

        resposta = self.client.get(URL_PDF)

        self.assertEqual(resposta.status_code, status.HTTP_200_OK)
        self.assertEqual(resposta["Content-Type"], "application/pdf")
        self.assertTrue(resposta.content.startswith(b"%PDF-"))
        self.assertIn("faturas-pendentes-2026-09-09.pdf", resposta["Content-Disposition"])

    def test_pdf_vazio_e_gerado(self, buscar):
        buscar.return_value = resposta_fedhub([])

        resposta = self.client.get(URL_PDF)

        self.assertEqual(resposta.status_code, status.HTTP_200_OK)
        self.assertTrue(resposta.content.startswith(b"%PDF-"))

    def test_pdf_longo_pagina_e_numera(self, buscar):
        buscar.return_value = resposta_fedhub([linha(documento=f"{i:010d}") for i in range(140)])

        resposta = self.client.get(URL_PDF)

        self.assertEqual(resposta.status_code, status.HTTP_200_OK)
        # Um PDF de 140 documentos em dois renglões não cabe numa página.
        self.assertGreater(resposta.content.count(b"/Type /Page\n") + resposta.content.count(b"/Type /Page "), 1)


class DocumentosTests(APITestCase):
    """Funções puras: linhas da tabela, filtros descritos, moeda."""

    def test_dois_renglioes_por_documento_com_produto_obs_e_administradora(self):
        dados, estilos = documentos.linhas_da_tabela([linha(obs="Aguardando síndico")])

        self.assertEqual(len(dados), 4)  # 2 de cabeçalho + 2 do documento
        self.assertIn("PROGRAMA DE GESTÃO", dados[3][1].text)
        self.assertIn("Aguardando síndico", dados[3][1].text)
        self.assertEqual(dados[3][4].text, "4/1")
        self.assertEqual(dados[3][5].text, "ADMINISTRADORA EXEMPLO LTDA")
        self.assertEqual(dados[2][5].text, "R$ 190,80")
        self.assertTrue(any(e[0] == "LINEBELOW" and e[1] == (0, 3) for e in estilos))

    def test_descricao_dos_filtros_e_moeda(self):
        texto = documentos.descrever_filtros({"situacao": "a_vencer", "classificacao": "deposito_cc", "administradora": "5"})

        self.assertEqual(texto, ["Situação: A vencer", "Depósito em conta corrente", "Administradora: 5"])
        self.assertEqual(documentos.moeda_br(45581.23), "R$ 45.581,23")
        self.assertEqual(documentos.moeda_br(None), "R$ 0,00")

    def test_filtrar_parametros_ignora_vazios_e_fora_do_contrato(self):
        from django.http import QueryDict

        q = QueryDict("fatura=10&situacao=&produtor=1&vencimento_ini=2026-01-01&classificacao=null")

        self.assertEqual(filtrar_parametros(q), {"fatura": "10", "vencimento_ini": "2026-01-01"})


class ServicoTests(APITestCase):
    def test_timeout_e_conexao_viram_excecoes_proprias(self):
        servico = RelatoriosFedhubService(base_url="http://fedhub.test")
        with patch("fedhub.services.relatorios_service.requests.get", side_effect=requests.Timeout()):
            with self.assertRaises(FedHubDemorou):
                servico.faturas_pendentes({})
        with patch("fedhub.services.relatorios_service.requests.get", side_effect=requests.ConnectionError()):
            with self.assertRaises(FedHubIndisponivel):
                servico.faturas_pendentes({})

    def test_status_diferente_de_200_vira_recusa_com_detalhe(self):
        servico = RelatoriosFedhubService(base_url="http://fedhub.test")

        class Resp:
            status_code = 400

            def json(self):
                return {"detail": "situacao inválida"}

        with patch("fedhub.services.relatorios_service.requests.get", return_value=Resp()):
            with self.assertRaises(FedHubRecusou) as ctx:
                servico.faturas_pendentes({"situacao": "x"})
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertEqual(ctx.exception.detalhe, {"detail": "situacao inválida"})
