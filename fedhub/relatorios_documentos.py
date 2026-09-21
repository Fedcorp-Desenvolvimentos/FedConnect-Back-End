"""Planilha e PDF do relatório de faturas pendentes (RF-FAT-002, RF-FAT-003).

Funções puras sobre a resposta do FedHub: recebem `dados` (`data`, `totais`,
`filtros`, `referencia`) e devolvem bytes.

O PDF nasceu como cópia do relatório do legado (PA-014) e foi redesenhado em
2026-09-21 (PA-036) com a linguagem visual do voucher de comissão: faixa azul
institucional no topo, cartões de metadados, tabela zebrada com cabeçalho
azul e uma linha por documento. O conteúdo é o mesmo, menos as duas colunas
de pagamento do legado — num relatório de pendentes elas são sempre vazias.
"""
from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

TITULO = "RELATÓRIO DE FATURAS PENDENTES"

ROTULO_SITUACAO = {"vencidas": "Vencidas", "a_vencer": "A vencer", "todas": "Todas"}
ROTULO_CLASSIFICACAO = {
    "todas": "Todas as faturas",
    "com_obs": "Somente com OBS preenchida",
    "sem_obs_sem_deposito": "OBS não preenchida e sem depósito em C/C",
    "deposito_cc": "Depósito em conta corrente",
}

COLUNAS_PLANILHA = [
    ("fatura", "Fatura"),
    ("documento", "Documento"),
    ("sacado", "Sacado"),
    ("produto", "Produto"),
    ("obs", "OBS"),
    ("vigencia", "Vigência"),
    ("vencimento", "Vencimento"),
    ("dias_atraso", "Dias em atraso"),
    ("valor", "Valor do documento (R$)"),
    ("parcela", "Parcela"),
    ("administradora_nome", "Administradora"),
    ("seguradora_nome", "Seguradora"),
    ("cedente", "Cedente"),
    ("apolice", "Apólice"),
    ("nosso_numero", "Nosso número"),
    ("deposito_cc", "Depósito em C/C"),
]


def _data_br(valor) -> str:
    if not valor:
        return ""
    if isinstance(valor, (date, datetime)):
        return valor.strftime("%d/%m/%Y")
    try:
        return datetime.fromisoformat(str(valor)[:10]).strftime("%d/%m/%Y")
    except ValueError:
        return str(valor)


def _data_iso_para_date(valor):
    if not valor:
        return None
    try:
        return datetime.fromisoformat(str(valor)[:10]).date()
    except ValueError:
        return None


def moeda_br(valor) -> str:
    numero = float(valor or 0)
    inteiro, decimal = f"{numero:,.2f}".split(".")
    return f"R$ {inteiro.replace(',', '.')},{decimal}"


def nome_arquivo(extensao: str, referencia: str = "") -> str:
    dia = referencia or date.today().isoformat()
    return f"faturas-pendentes-{dia}.{extensao}"


def descrever_filtros(filtros: Dict[str, Any]) -> List[str]:
    """Filtros aplicados em texto, para o cabeçalho do PDF e a aba da planilha."""
    partes = []
    if filtros.get("situacao"):
        partes.append(f"Situação: {ROTULO_SITUACAO.get(filtros['situacao'], filtros['situacao'])}")
    if filtros.get("classificacao"):
        partes.append(ROTULO_CLASSIFICACAO.get(filtros["classificacao"], filtros["classificacao"]))
    for chave, rotulo in (("fatura", "Fatura"), ("administradora", "Administradora"), ("seguradora", "Seguradora"),
                          ("cedente", "Cedente"), ("apolice", "Apólice")):
        if filtros.get(chave):
            partes.append(f"{rotulo}: {filtros[chave]}")
    if filtros.get("vigencia_ini"):
        partes.append(f"Vigência a partir de {_data_br(filtros['vigencia_ini'])}")
    if filtros.get("ordenar_por"):
        partes.append(f"Ordenado por {filtros['ordenar_por']}")
    return partes


# ---------------------------------------------------------------------------
# Planilha
# ---------------------------------------------------------------------------

def gerar_planilha(dados: Dict[str, Any]) -> bytes:
    """Uma linha por documento, valores como número, datas como data, linha de totais."""
    linhas = dados.get("data", [])
    totais = dados.get("totais", {})
    wb = Workbook()
    ws = wb.active
    ws.title = "Faturas pendentes"

    cabecalho_fill = PatternFill("solid", fgColor="0F3D5D")
    for indice, (_, rotulo) in enumerate(COLUNAS_PLANILHA, start=1):
        celula = ws.cell(row=1, column=indice, value=rotulo)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = cabecalho_fill
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for r, linha in enumerate(linhas, start=2):
        for c, (chave, _) in enumerate(COLUNAS_PLANILHA, start=1):
            valor = linha.get(chave)
            if chave == "vencimento":
                valor = _data_iso_para_date(valor)
                celula = ws.cell(row=r, column=c, value=valor)
                celula.number_format = "DD/MM/YYYY"
                continue
            if chave == "valor":
                celula = ws.cell(row=r, column=c, value=float(valor or 0))
                celula.number_format = '"R$" #,##0.00'
                continue
            ws.cell(row=r, column=c, value=valor)

    linha_total = len(linhas) + 2
    ws.cell(row=linha_total, column=1, value=f"{totais.get('documentos', len(linhas))} documento(s)").font = Font(bold=True)
    ws.cell(row=linha_total, column=2, value=f"{totais.get('faturas', 0)} fatura(s)").font = Font(bold=True)
    ws.cell(row=linha_total, column=8, value="TOTAL GERAL").font = Font(bold=True)
    celula_total = ws.cell(row=linha_total, column=9, value=float(totais.get("valor_total", 0)))
    celula_total.font = Font(bold=True)
    celula_total.number_format = '"R$" #,##0.00'

    larguras = [10, 14, 38, 36, 24, 10, 12, 10, 18, 10, 36, 24, 10, 12, 16, 18]
    for indice, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = largura
    ws.freeze_panes = "A2"

    aba = wb.create_sheet("Filtros")
    aba.append(["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
    aba.append(["Referência", _data_br(dados.get("referencia"))])
    for texto in descrever_filtros(dados.get("filtros", {})):
        aba.append([texto])
    aba.column_dimensions["A"].width = 60

    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

# Paleta do voucher (`voucher_controller.py`) trazida para o ReportLab.
AZUL = colors.HexColor("#0F3D5D")          # institucional: faixas e títulos
AZUL_CLARO = colors.HexColor("#1B5478")    # segundo tom da faixa do topo
CINZA = colors.HexColor("#64748B")
CINZA_CLARO = colors.HexColor("#94A3B8")
PONTILHADO = colors.HexColor("#94A3B8")
LINHA = colors.HexColor("#E2E8F0")
ZEBRA = colors.HexColor("#F7FAFC")
VERDE = colors.HexColor("#1F7A4D")         # valores, como no voucher

_estilo_base = ParagraphStyle("base", fontName="Helvetica", fontSize=7.4, leading=9)
_estilo_negrito = ParagraphStyle("negrito", parent=_estilo_base, fontName="Helvetica-Bold")
_estilo_direita = ParagraphStyle("direita", parent=_estilo_base, alignment=TA_RIGHT)
_estilo_cabecalho = ParagraphStyle(
    "cab", parent=_estilo_negrito, fontSize=6.4, leading=7.6,
    textColor=colors.white, alignment=TA_LEFT,
)
_estilo_cabecalho_dir = ParagraphStyle("cabdir", parent=_estilo_cabecalho, alignment=TA_RIGHT)
_estilo_cinza = ParagraphStyle("cinza", parent=_estilo_base, fontSize=6.8, leading=8.2, textColor=CINZA)
_estilo_mono = ParagraphStyle("mono", parent=_estilo_base, fontName="Courier", fontSize=7.2)
_estilo_valor = ParagraphStyle(
    "valor", parent=_estilo_base, fontName="Courier-Bold", fontSize=7.6,
    textColor=VERDE, alignment=TA_RIGHT,
)
_estilo_rotulo = ParagraphStyle(
    "rotulo", parent=_estilo_base, fontName="Helvetica-Bold", fontSize=5.9,
    leading=7, textColor=CINZA,
)
_estilo_dado = ParagraphStyle("dado", parent=_estilo_negrito, fontSize=8.6, leading=10)


class _CanvasNumerado(rl_canvas.Canvas):
    """Faixa azul no topo e rodapé com "Página N de M" (M só é conhecido no fim)."""

    def __init__(self, *args, cabecalho=None, **kwargs):
        super().__init__(*args, **kwargs)
        self._paginas = []
        self._cabecalho = cabecalho or {}

    def showPage(self):
        self._paginas.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._paginas)
        for estado in self._paginas:
            self.__dict__.update(estado)
            self._desenhar_cabecalho(total)
            super().showPage()
        super().save()

    def _desenhar_cabecalho(self, total):
        """Faixa azul com a marca e o título, e o rodapé da página.

        Desenhar no canvas, e não como elemento do fluxo, é o que faz a faixa
        sangrar até a borda e se repetir em toda página — como o cabeçalho do
        voucher de comissão.
        """
        largura, altura = A4
        faixa = 2.35 * cm

        # Degradê em faixas finas: o ReportLab não tem gradiente, e dois
        # retângulos deixavam uma emenda visível no meio do cabeçalho.
        passos = 48
        for i in range(passos):
            self.setFillColor(AZUL.clone() if i == 0 else _mistura(AZUL, AZUL_CLARO, i / (passos - 1)))
            self.rect(largura * i / passos, altura - faixa,
                      largura / passos + 0.6, faixa, stroke=0, fill=1)

        self.setFillColor(colors.white)
        self.setFont("Helvetica-Bold", 7.5)
        self.drawString(1.2 * cm, altura - 0.95 * cm, "GRUPO FEDCORP")
        self.setFont("Helvetica-Bold", 15)
        self.drawString(1.2 * cm, altura - 1.75 * cm, "Faturas pendentes")

        self.setFont("Helvetica", 7.5)
        self.drawRightString(largura - 1.2 * cm, altura - 0.95 * cm,
                             f"Emitido em {self._cabecalho.get('data_geracao', '')}")
        self.drawRightString(largura - 1.2 * cm, altura - 1.75 * cm,
                             self._cabecalho.get("linha_periodo", ""))

        self.setFillColor(CINZA_CLARO)
        self.setFont("Helvetica", 6.4)
        self.drawString(1.2 * cm, 0.75 * cm, self._cabecalho.get("linha_filtros", "")[:150])
        self.drawRightString(largura - 1.2 * cm, 0.75 * cm, f"Página {self._pageNumber} de {total}")
        self.setStrokeColor(LINHA)
        self.setLineWidth(0.4)
        self.line(1.2 * cm, 1.05 * cm, largura - 1.2 * cm, 1.05 * cm)


def _mistura(inicio, fim, t: float):
    """Cor intermediária entre duas, para simular degradê."""
    return colors.Color(
        inicio.red + (fim.red - inicio.red) * t,
        inicio.green + (fim.green - inicio.green) * t,
        inicio.blue + (fim.blue - inicio.blue) * t,
    )


def _escapar(texto) -> str:
    return (texto or "").replace("&", "&amp;").replace("<", "&lt;")


def _paragrafo(texto, estilo=_estilo_base):
    """Texto simples: escapa tudo, porque vem do ERP."""
    return Paragraph(_escapar(texto), estilo)


def _paragrafo_duplo(principal, secundario, estilo, cor="#64748B", tamanho="6.6"):
    """Duas linhas na mesma célula: o dado e, abaixo, o complemento em cinza.

    É o que substitui o segundo renglão do legado. As duas partes são
    escapadas aqui; as tags são nossas.
    """
    if not secundario:
        return Paragraph(_escapar(principal), estilo)
    return Paragraph(
        f"{_escapar(principal)}<br/><font color='{cor}' size='{tamanho}'>{_escapar(secundario)}</font>",
        estilo,
    )


def linhas_da_tabela(linhas: List[Dict[str, Any]]):
    """Cabeçalho + **uma** linha por documento, e os comandos de estilo.

    O legado usava dois renglões por documento porque espremia dez colunas,
    duas delas sempre vazias (data e valor de pagamento — em relatório de
    pendentes não há pagamento). Sem elas cabe uma linha só, com o que era o
    segundo renglão virando texto secundário na própria célula: produto e OBS
    sob o sacado, dias em atraso sob o vencimento.
    """
    dados = [[
        _paragrafo("FATURA", _estilo_cabecalho),
        _paragrafo("DOCUMENTO", _estilo_cabecalho),
        _paragrafo("SACADO / PRODUTO", _estilo_cabecalho),
        _paragrafo("ADMINISTRADORA", _estilo_cabecalho),
        _paragrafo("VIGÊNCIA", _estilo_cabecalho),
        _paragrafo("VENCIMENTO", _estilo_cabecalho),
        _paragrafo("PARC", _estilo_cabecalho),
        _paragrafo("VALOR", _estilo_cabecalho_dir),
    ]]
    estilos = [
        ("BACKGROUND", (0, 0), (-1, 0), AZUL),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, 0), 6),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
    ]

    for indice, linha in enumerate(linhas):
        r = indice + 1
        sacado = linha.get("sacado") or ""
        produto_obs = linha.get("produto") or ""
        if linha.get("obs"):
            produto_obs = f"{produto_obs} — {linha['obs']}" if produto_obs else linha["obs"]
        vencimento = _data_br(linha.get("vencimento"))
        dias = linha.get("dias_atraso") or 0
        atraso = f"há {dias} dias" if dias else "no prazo"

        dados.append([
            _paragrafo(str(linha.get("fatura") or ""), _estilo_negrito),
            _paragrafo(linha.get("documento"), _estilo_mono),
            _paragrafo_duplo(sacado, produto_obs, _estilo_base),
            _paragrafo(linha.get("administradora_nome") or linha.get("administradora"), _estilo_cinza),
            _paragrafo(linha.get("vigencia"), _estilo_mono),
            _paragrafo_duplo(vencimento, atraso, _estilo_mono, cor="#94A3B8", tamanho="6.2"),
            _paragrafo(linha.get("parcela"), _estilo_mono),
            _paragrafo(moeda_br(linha.get("valor")), _estilo_valor),
        ])
        if indice % 2:
            estilos.append(("BACKGROUND", (0, r), (-1, r), ZEBRA))
        estilos.append(("LINEBELOW", (0, r), (-1, r), 0.4, LINHA))
    return dados, estilos


def _cartoes_resumo(dados: Dict[str, Any], largura_util: float) -> Table:
    """Faixa de metadados abaixo do título, no espírito dos cartões do voucher."""
    totais = dados.get("totais", {})
    filtros = dados.get("filtros", {})
    celulas = [
        ("SITUAÇÃO", ROTULO_SITUACAO.get(filtros.get("situacao", ""), "Todas")),
        ("DOCUMENTOS", str(totais.get("documentos", 0))),
        ("FATURAS", str(totais.get("faturas", 0))),
        ("TOTAL PENDENTE", moeda_br(totais.get("valor_total", 0))),
    ]
    tabela = Table(
        [[_paragrafo(r, _estilo_rotulo) for r, _ in celulas],
         [_paragrafo(v, _estilo_dado) for _, v in celulas]],
        colWidths=[largura_util / len(celulas)] * len(celulas),
    )
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7FAFC")),
        ("BOX", (0, 0), (-1, -1), 0.5, LINHA),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.white),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 7),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 0),
        ("TOPPADDING", (0, 1), (-1, 1), 1),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 7),
        ("TEXTCOLOR", (3, 1), (3, 1), AZUL),
    ]))
    return tabela


def gerar_pdf(dados: Dict[str, Any]) -> bytes:
    linhas = dados.get("data", [])
    totais = dados.get("totais", {})
    filtros = dados.get("filtros", {})
    periodo_ini = _data_br(filtros.get("vencimento_ini"))
    periodo_fim = _data_br(filtros.get("vencimento_fim"))
    if periodo_ini or periodo_fim:
        periodo = f"Vencimento {periodo_ini or '...'} a {periodo_fim or '...'}"
    else:
        periodo = "Todos os vencimentos"

    cabecalho = {
        "data_geracao": datetime.now().strftime("%d/%m/%Y"),
        "linha_periodo": periodo,
        "linha_filtros": " · ".join(descrever_filtros(filtros)),
    }

    saida = BytesIO()
    doc = SimpleDocTemplate(
        saida, pagesize=A4,
        leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=2.9 * cm, bottomMargin=1.4 * cm,
        title=TITULO, author="FedConnect",
    )
    largura_util = A4[0] - doc.leftMargin - doc.rightMargin
    # Sem as colunas de pagamento sobra largura para o sacado e a administradora.
    # Medidas conferidas contra o texto mais largo de cada coluna: documento e
    # vencimento têm 10 caracteres em Courier 7,2 (~43 pt) e não podem quebrar.
    proporcoes = [0.08, 0.115, 0.27, 0.155, 0.085, 0.115, 0.06, 0.12]
    larguras = [largura_util * p for p in proporcoes]

    dados_tabela, estilos = linhas_da_tabela(linhas)
    tabela = Table(dados_tabela, colWidths=larguras, repeatRows=1)
    tabela.setStyle(TableStyle(estilos))

    rodape = Table(
        [[
            _paragrafo(
                f"{totais.get('faturas', 0)} fatura(s) · {totais.get('documentos', len(linhas))} documento(s)",
                ParagraphStyle("rf", parent=_estilo_base, fontSize=8, textColor=colors.white),
            ),
            _paragrafo("TOTAL GERAL", ParagraphStyle(
                "rt", parent=_estilo_negrito, fontSize=8, textColor=colors.white, alignment=TA_RIGHT)),
            _paragrafo(moeda_br(totais.get("valor_total", 0)), ParagraphStyle(
                "rv", parent=_estilo_negrito, fontName="Courier-Bold", fontSize=11,
                textColor=colors.white, alignment=TA_RIGHT)),
        ]],
        colWidths=[largura_util * 0.55, largura_util * 0.2, largura_util * 0.25],
    )
    rodape.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), AZUL),
        ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, 0), 9),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 9),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
    ]))

    elementos = [_cartoes_resumo(dados, largura_util), Spacer(1, 0.45 * cm)]
    if linhas:
        elementos += [tabela, Spacer(1, 0.35 * cm), rodape]
    else:
        # Sem linhas, uma tabela só com cabeçalho e uma faixa de total zerada
        # ficam sem sentido: entra um aviso centralizado no lugar das duas.
        vazio = Table(
            [[_paragrafo("Nenhuma fatura pendente para os filtros informados.",
                         ParagraphStyle("vz", parent=_estilo_base, fontSize=9,
                                        textColor=CINZA, alignment=TA_CENTER))]],
            colWidths=[largura_util],
        )
        vazio.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.5, LINHA),
            ("BACKGROUND", (0, 0), (-1, -1), ZEBRA),
            ("TOPPADDING", (0, 0), (-1, -1), 26),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 26),
        ]))
        elementos.append(vazio)
    doc.build(elementos, canvasmaker=lambda *a, **k: _CanvasNumerado(*a, cabecalho=cabecalho, **k))
    return saida.getvalue()
